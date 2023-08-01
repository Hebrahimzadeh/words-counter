    #!/usr/bin/python

import sys
import docx
import os
import re
from farsi_tools import standardize_persian_text
import openpyxl


def doc_files(path):
    files = [filename for filename in os.listdir(path) if filename.endswith('.docx')]
    return files

def get_doc_texts(doc_file):
    doc = docx.Document(doc_file)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(standardize_persian_text(para.text))
    return fullText

def text_words(text):
    return re.findall(r'\w+', '\n'.join(text))

if __name__ == '__main__':
    file_path = './files'
    files = doc_files(file_path)
    xfile = openpyxl.Workbook()
    xfile_writer = xfile.active
    xfile_writer.title = 'words'
    
    for file in files:
        xfile_writer.append((' کلمات ' + file,))            
        file_text = get_doc_texts(file_path + '/' + file)
        for word in text_words(file_text):
            xfile_writer.append((word,))
    
    xfile.save('words.xlsx')
        
    

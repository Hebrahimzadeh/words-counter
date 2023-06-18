    #!/usr/bin/python

import sys
import docx
import re
import os
from farsi_tools import standardize_persian_text
import openpyxl


def doc_files(path):
    files = [filename for filename in os.listdir(path) if filename.rstrip('.docx')]
    return files

def getTags(file_path = 'tags.txt'):
    with open(file_path,  'r', encoding='utf8') as f:
        lines = [line.strip() for line in f.readlines()]
    return lines

def get_doc_text(doc_file):
    doc = docx.Document(doc_file)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(standardize_persian_text(para.text))
    return '\n'.join(fullText)

def count_word(word, text):
    return text.count(word)


if __name__ == '__main__':
    file_path = './files'
    files = doc_files(file_path)
    tags = getTags()
    total_tags_count = {}
    total_tag_count_in_files={}
    xfile_tags = openpyxl.Workbook()
    xfile_tags_writer = xfile_tags.active
    xfile_tags_writer.title = 'file and tags';
    
    xfile_tags_writer.append(('نام فایل', 'تگ‌ها', 'تعداد تگ‌ها'))            
    for file in files:
        file_tags_count = {}
        file_tags = {}
        file_text = get_doc_text(file_path + '/' + file)
        for tag in tags:
            if tag not in total_tags_count:
                total_tags_count[tag] = 0
            if tag not in total_tag_count_in_files:
                total_tag_count_in_files[tag] = 0
                
            counts = count_word(tag, file_text)
        
            if counts > 0:
                total_tags_count[tag] += counts
                file_tags[tag] = counts
                total_tag_count_in_files[tag] += 1
                
        file_tags_count[file] = { "tags": ','.join([tag for tag in file_tags]), "tags_count": len(file_tags), "tag_total_count": sum(file_tags.values()) }
        for file,data in file_tags_count.items(): 
            xfile_tags_writer.append((file, data["tags"], data["tags_count"]))
            
    xfile_tags.create_sheet('all')
    xfile_total_tags_writer = xfile_tags['all']
    
    xfile_total_tags_writer.append(('تگ', 'تعداد کل استفاده‌ها', 'تعداد فایل‌های استفاده شده'))
    for tag, count in total_tags_count.items():
        xfile_total_tags_writer.append((tag, count, total_tag_count_in_files[tag]))
    
    xfile_tags.save('file_tags.xlsx')
        
    
